"""AI 简历生成接口 —— 本地 Ollama（专用模型 resume-writer，基于 Qwen3-8B）"""
import os
from io import BytesIO
import requests
from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from models import User, Achievement

ai_bp = Blueprint('ai', __name__, url_prefix='/api/ai')

OLLAMA_BASE_URL = os.getenv('OLLAMA_BASE_URL', 'http://localhost:11434')
OLLAMA_MODEL = os.getenv('OLLAMA_MODEL', 'resume-writer')

RESUME_SYSTEM_PROMPT = (
    "你是一名专业的简历撰写助手。\n"
    "你的唯一职责：根据用户提供的个人成果信息，生成一份结构清晰、排版整洁的中文简历。\n"
    "规则：\n"
    "1. 用户提供的成果按'获奖经历、证书认证、项目经历、学术成果、荣誉称号、社会实践'等归类写入对应章节。\n"
    "2. 生成简历时使用以下结构分节输出：基本信息、教育经历、获奖经历、项目经历、技能特长、自我评价。\n"
    "3. 获奖/证书/项目等条目保留原始信息（名称、级别、时间），并适当润色为专业表述，突出成果与量化数据。\n"
    "4. 用户成果中未提供的信息该节写'待补充'，不要编造。\n"
    "5. 若有用户的补充说明，按其意图调整（如突出某方面、目标岗位等）。\n"
    "6. 只输出简历正文，不要额外解释。"
)

# Word 导出的章节标题
RESUME_SECTIONS = ['基本信息', '教育经历', '获奖经历', '项目经历', '技能特长', '自我评价']


def _set_cn_font(run, size=12, bold=False, color=None):
    """设置中文字体（微软雅黑），避免 Word 中文乱码/默认字体"""
    run.font.name = 'Microsoft YaHei'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def build_resume_docx(resume_text):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(11)

    lines = [l.strip() for l in (resume_text or '').splitlines() if l.strip()]
    for line in lines:
        if any(line.startswith(s) for s in RESUME_SECTIONS):
            p = doc.add_paragraph()
            _set_cn_font(p.add_run(line), size=15, bold=True, color=RGBColor(0x2F, 0x80, 0xED))
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(2)
    return doc


def _format_achievement(ach):
    """把单条成果格式化为简历素材文本"""
    parts = [ach.title]
    if ach.level:
        parts.append(f'（{ach.level}）')
    if ach.achieved_date:
        parts.append(f'，时间 {ach.achieved_date.strftime("%Y-%m")}')
    text = ''.join(parts)
    if ach.description:
        text += f'；描述：{ach.description}'
    if ach.keywords:
        text += f'；关键词：{ach.keywords}'
    return text


@ai_bp.route('/resume', methods=['POST'])
@jwt_required()
def generate_resume():
    """根据学生勾选的成果记录调用本地模型生成简历"""
    data = request.get_json(silent=True) or {}
    achievement_ids = data.get('achievement_ids') or []
    extra = (data.get('extra') or '').strip()

    if not achievement_ids:
        return jsonify({'error': '请先勾选你的成果记录'}), 400

    current_user_id = int(get_jwt_identity())
    user = User.query.get(current_user_id)
    if not user or user.role != 'student' or not user.student:
        return jsonify({'error': '仅学生可使用简历生成'}), 403
    student = user.student

    query = Achievement.query.filter(
        Achievement.student_id == student.id,
        Achievement.id.in_(achievement_ids)
    )
    achievements = query.order_by(Achievement.achieved_date.desc()).all()

    if not achievements:
        return jsonify({'error': '未找到勾选的成果记录'}), 400

    # 按一级分类分组，方便模型归类进简历章节
    groups = {}
    for ach in achievements:
        groups.setdefault(ach.main_category or '其他', []).append(_format_achievement(ach))

    lines = [f'{cat}：' + '；'.join(items) for cat, items in groups.items()]
    achievement_text = '\n'.join(lines)

    user_content = f'学生的成果记录如下：\n{achievement_text}\n'
    if extra:
        user_content += f'\n补充要求：{extra}\n'

    try:
        resp = requests.post(
            f'{OLLAMA_BASE_URL}/v1/chat/completions',
            json={
                'model': OLLAMA_MODEL,
                'messages': [
                    {'role': 'system', 'content': RESUME_SYSTEM_PROMPT},
                    {'role': 'user', 'content': user_content}
                ],
                'temperature': 0.7,
                'max_tokens': 4000,
                'stream': False
            },
            timeout=180
        )
        if resp.status_code != 200:
            if resp.status_code == 404:
                return jsonify({'error': f'模型 {OLLAMA_MODEL} 不存在，请先运行：ollama create resume-writer -f Modelfile'}), 502
            return jsonify({'error': f'Ollama 服务异常: {resp.status_code}'}), 502
        content = resp.json()['choices'][0]['message']['content']
        return jsonify({'resume': content})
    except requests.exceptions.ConnectionError:
        return jsonify({'error': '无法连接 Ollama 服务，请确认已运行 ollama serve'}), 502
    except Exception as e:
        return jsonify({'error': f'生成失败: {str(e)}'}), 500


@ai_bp.route('/resume/export', methods=['POST'])
@jwt_required()
def export_resume():
    """把已生成的简历文本排版为 Word 文档并下载"""
    data = request.get_json(silent=True) or {}
    resume_text = (data.get('resume') or '').strip()
    if not resume_text:
        return jsonify({'error': '暂无简历内容可导出'}), 400

    doc = build_resume_docx(resume_text)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name='我的简历.docx',
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
